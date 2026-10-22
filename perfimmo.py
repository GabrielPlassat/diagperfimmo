"""
Moteur d'extraction Diag Perf'Immo -> Excel, via l'API Albert (compatible OpenAI).

Étapes :
  1. lire_pdf()            : PDF -> texte Markdown page par page (tableaux conservés)
  2. lire_indicateurs()    : liste des indicateurs = colonne A du modèle Excel
  3. extraire_rapport()    : appels Albert par groupe d'indicateurs -> JSON
  4. verifier()            : contrôles automatiques (valeur présente dans le PDF, cohérences)
  5. ecrire_excel()        : remplit le modèle, couleur = niveau de confiance
"""
from __future__ import annotations

import csv
import io
import json
import re
import time
import unicodedata
from concurrent.futures import ThreadPoolExecutor, as_completed
from copy import copy
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ALBERT_BASE_URL = "https://albert.api.etalab.gouv.fr/v1"
DOSSIER = Path(__file__).parent
MODELE_PAR_DEFAUT = DOSSIER / "modele" / "modele_indicateurs_perfimmo_enrichis.xlsx"
CONSIGNES_CSV = DOSSIER / "consignes.csv"

# Modèles préférés, dans l'ordre (identifiants ou alias Albert)
MODELES_PREFERES = [
    "gpt-oss-120b", "openweight-large",
    "mistral-small-3-2-24b-instruct-2506", "openweight-medium",
    "gemma-4-31b-it", "deepseek-v4-flash",
    "ministral-3-8b-instruct-2512", "openweight-small",
]

CONFIANCES = ["haute", "moyenne", "faible", "non trouvé", "validé"]
COULEURS = {  # fond de cellule Excel
    "haute": "C6EFCE",      # vert
    "moyenne": "FFEB9C",    # jaune
    "faible": "FFC7CE",     # rouge
    "non trouvé": "EDEDED", # gris
    "validé": "BDD7EE",     # bleu : corrigé / validé à la main
}


# ---------------------------------------------------------------------------
# 1. Lecture du PDF
# ---------------------------------------------------------------------------
def lire_pdf(pdf_bytes: bytes) -> list[str]:
    """Retourne une liste de textes (un par page). Essaie d'abord la conversion
    Markdown (meilleure pour les tableaux), sinon texte brut."""
    import pymupdf

    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    try:
        import pymupdf4llm

        try:
            chunks = pymupdf4llm.to_markdown(
                doc, page_chunks=True, use_ocr=False, header=False, footer=False,
                show_progress=False,
            )
        except TypeError:  # anciennes versions de pymupdf4llm
            chunks = pymupdf4llm.to_markdown(doc, page_chunks=True, show_progress=False)
        pages = [c.get("text", "") for c in chunks]
        if len(pages) == len(doc) and sum(len(p) for p in pages) > 500:
            return [_nettoyer(p) for p in pages]
    except Exception:
        pass
    return [_nettoyer(p.get_text()) for p in doc]


def _nettoyer(t: str) -> str:
    t = re.sub(r"<!-- (Start|End) of picture text -->", "", t)
    t = t.replace("<br>", " / ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


def image_page(pdf_bytes: bytes, num_page: int, zoom: float = 1.6) -> bytes:
    """PNG d'une page (numérotée à partir de 1) pour la relecture humaine."""
    import pymupdf

    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    num_page = max(1, min(num_page, len(doc)))
    pix = doc[num_page - 1].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom))
    return pix.tobytes("png")


# ---------------------------------------------------------------------------
# 2. Indicateurs (lus dans le modèle Excel) et consignes
# ---------------------------------------------------------------------------
@dataclass
class Indicateur:
    ligne: int           # ligne dans le modèle Excel d'origine
    libelle: str
    groupe: str = ""
    consigne: str = ""

    @property
    def unite(self) -> str:
        m = re.findall(r"\(([^()]*)\)\s*$", self.libelle.strip())
        return m[-1] if m else ""


GROUPES = {
    "identification": "Identification de l'entreprise, du site et du projet (dates, adresse, surface, usage, occupation)",
    "reference": "Consommations actuelles et de l'année de référence, objectifs du décret tertiaire (Cref, Cabs, Crelat), répartition par énergie et par usage, coût",
    "deperditions": "Déperditions thermiques du bâtiment (en kW)",
    "gains_rapides": "Scénario / plan d'actions à gains rapides (actions immédiates, faible investissement)",
    "2030": "Scénario visant l'échéance 2030 du décret tertiaire (souvent « Scénario 1 »)",
    "2040": "Scénario visant l'échéance 2040 du décret tertiaire (souvent « Scénario 2 »)",
    "2050": "Scénario visant l'échéance 2050 du décret tertiaire (souvent « Scénario 3 »)",
}

# Repères qui font changer de groupe en parcourant le modèle de haut en bas
_DEBUTS_GROUPES = [
    (r"consommation ann[ée]e courante|^mois\s*-\s*r[ée]f[ée]rence|^ann[ée]e\s*-\s*r[ée]f[ée]rence|^cabs\s*-\s*r[ée]f", "reference"),
    (r"^d[ée]perditions", "deperditions"),
    (r"^(nombre|nb) pr[ée]conisations\s*-\s*gains rapides|^gains rapides$", "gains_rapides"),
    (r"^(nombre|nb) pr[ée]conisations\s*-\s*2030|^ann[ée]e\s*-\s*2030", "2030"),
    (r"^(nombre|nb) pr[ée]conisations\s*-\s*2040|^ann[ée]e\s*-\s*2040", "2040"),
    (r"^(nombre|nb) pr[ée]conisations\s*-\s*2050|^ann[ée]e\s*-\s*2050", "2050"),
]


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return re.sub(r"\s+", " ", s).strip().lower()


def lire_consignes(chemin: Path = CONSIGNES_CSV) -> dict[str, str]:
    if not chemin.exists():
        return {}
    with open(chemin, encoding="utf-8-sig", newline="") as f:
        return {_norm(r["indicateur"]): r.get("consigne", "").strip()
                for r in csv.DictReader(f, delimiter=";") if r.get("indicateur")}


def lire_indicateurs(modele_bytes: bytes, consignes: dict[str, str] | None = None
                     ) -> tuple[list[Indicateur], str]:
    """Lit la colonne A de la 1re feuille non vide du modèle."""
    consignes = consignes if consignes is not None else lire_consignes()
    wb = openpyxl.load_workbook(io.BytesIO(modele_bytes))
    ws = next((w for w in wb.worksheets if w.max_row > 5), wb.worksheets[0])
    inds, groupe = [], "identification"
    for r in range(1, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if v is None or not str(v).strip():
            continue
        lib = str(v).strip()
        for motif, g in _DEBUTS_GROUPES:
            if re.search(motif, _norm(lib)):
                groupe = g
                break
        inds.append(Indicateur(r, lib, groupe, consignes.get(_norm(lib), "")))
    return inds, ws.title


# ---------------------------------------------------------------------------
# 3. Appel à Albert
# ---------------------------------------------------------------------------
PROMPT_SYSTEME = """Tu es un ingénieur thermicien expert du décret tertiaire (dispositif Éco Énergie Tertiaire, DEET) et des audits « Diag Perf'Immo » (Bpifrance / ADEME).
Ta mission : extraire des valeurs précises d'un rapport d'audit énergétique pour remplir un tableau d'indicateurs.

Glossaire :
- Créf / Cref : consommation de l'année de référence (kWhef/m²), éventuellement ajustée/modulée.
- Cabs : objectif en valeur absolue (kWhef/m²) fixé par arrêté selon la catégorie d'activité ; souvent « NC » s'il n'est pas publié.
- Crelat : objectif en valeur relative = 60 % (2030), 50 % (2040), 40 % (2050) de Créf. « avec modulation » = Crelat modulé.
- SUB : surface utile brute ; SDP : surface de plancher ; SHON ; EF = énergie finale ; EP = énergie primaire.
- CEE : certificats d'économies d'énergie ; TRB : temps de retour brut ; « prix constants / fixes » vs « prix croissants ».
- Scénarios : le rapport présente en général un scénario par échéance (2030, 2040, 2050), parfois numérotés 1, 2, 3.

Règles strictes :
1. N'invente JAMAIS une valeur. Si l'information est absente, anonymisée ([adresse], [Prénom Nom], xxxxx…) ou notée « NC », mets "valeur": null et explique dans "commentaire".
2. Les nombres sont rendus en nombre JSON (point décimal, sans espace ni unité) : 634302 et non "634 302".
3. Respecte l'unité demandée entre parenthèses dans le libellé. Si le rapport donne une autre unité (MWh, €, kWh total…), convertis (ex. € -> k€, kWh -> kWh/m² en divisant par la surface indiquée dans le rapport), mets "confiance": "moyenne" et décris le calcul dans "commentaire".
4. "page" = numéro de page du PDF (repère « === PAGE n === ») où se trouve la valeur.
5. "extrait" = courte citation exacte du rapport (≤ 25 mots) qui justifie la valeur.
6. "confiance" :
   - "haute" : valeur lue telle quelle, sans ambiguïté, correspondant exactement à l'indicateur ;
   - "moyenne" : conversion, calcul, ou correspondance approximative avec l'indicateur ;
   - "faible" : déduction incertaine, plusieurs valeurs possibles, tableau difficile à lire.
7. Dates au format JJ/MM/AAAA (ou MM/AAAA, AAAA si seule l'année est connue). Oui/non en "oui" ou "non".
8. Réponds UNIQUEMENT avec un objet JSON valide, sans texte autour."""

FORMAT_REPONSE = """{"resultats": [
  {"id": <id de l'indicateur>, "valeur": <nombre, texte ou null>, "page": <entier ou null>,
   "extrait": "<citation>", "confiance": "haute|moyenne|faible", "commentaire": "<optionnel>"}
]}"""


def construire_prompt(groupe: str, inds: list[Indicateur], pages: list[str],
                      pages_choisies: list[int] | None = None) -> str:
    lignes = []
    for i in inds:
        l = f'- id {i.ligne} : « {i.libelle} »'
        if i.consigne:
            l += f"  → consigne : {i.consigne}"
        lignes.append(l)
    idx = pages_choisies if pages_choisies else list(range(len(pages)))
    texte = "\n\n".join(f"=== PAGE {n + 1} ===\n{pages[n]}" for n in idx)
    return (
        f"Thème de cette série d'indicateurs : {GROUPES.get(groupe, groupe)}.\n\n"
        f"INDICATEURS À EXTRAIRE ({len(inds)}) :\n" + "\n".join(lignes) +
        f"\n\nFORMAT DE RÉPONSE ATTENDU (un élément par indicateur, même si valeur null) :\n{FORMAT_REPONSE}\n\n"
        f"RAPPORT D'AUDIT ({len(idx)} pages sur {len(pages)}) :\n\n{texte}"
    )


_MOTS_GROUPES = {
    "identification": "raison sociale siret siren adresse code postal ville effectif naf site zone climatique surface sub sdp bâtiment usage occupants horaires operat propriétaire locataire assujetti décret cadrage visite restitution date interlocuteurs",
    "reference": "consommation référence cref créf cabs crelat modulé modulation année kwhef/m² électricité gaz réseau chaleur biomasse photovoltaïque autoconsommé répartition usage chauffage climatisation ventilation éclairage ecs cuisson bureautique coût facture €",
    "deperditions": "déperditions murs menuiseries toiture plancher ponts thermiques ventilation infiltrations kw",
    "gains_rapides": "gains rapides actions immédiates sobriété préconisations investissement cee gain temps de retour trb scénario 0 photovoltaïque",
    "2030": "scénario 1 2030 préconisations cabs crelat investissement cee gain temps de retour trb ratio consommations synthèse scénarios",
    "2040": "scénario 2 2040 préconisations cabs crelat investissement cee gain temps de retour trb ratio consommations synthèse scénarios",
    "2050": "scénario 3 2050 préconisations cabs crelat investissement cee gain temps de retour trb ratio consommations synthèse scénarios performance équipements bacs photovoltaïque",
}


def pages_pertinentes(groupe: str, inds: list[Indicateur], pages: list[str],
                      max_pages: int = 18) -> list[int]:
    """Sélection des pages les plus pertinentes par mots-clés (mode économe)."""
    mots = set(_norm(_MOTS_GROUPES.get(groupe, "")).split())
    for i in inds:
        mots |= {m for m in re.findall(r"[a-z0-9]{4,}", _norm(i.libelle))}
    scores = []
    for n, p in enumerate(pages):
        np_ = _norm(p)
        s = sum(np_.count(m) for m in mots) + 3 * np_.count("|")  / 50
        scores.append((s, n))
    choisies = sorted(n for _, n in sorted(scores, reverse=True)[:max_pages])
    if groupe == "identification":
        choisies = sorted(set(choisies) | set(range(min(3, len(pages)))))
    return choisies


def client_albert(api_key: str, base_url: str = ALBERT_BASE_URL):
    from openai import OpenAI
    return OpenAI(base_url=base_url, api_key=api_key, timeout=600, max_retries=0)


def lister_modeles(client) -> list[str]:
    try:
        ms = client.models.list().data
        ids = [m.id for m in ms if getattr(m, "type", "text-generation") in
               ("text-generation", "image-text-to-text", None)]
        ids = [i for i in ids if not re.search(r"ocr|whisper|embed|rerank|bge", i, re.I)]
    except Exception:
        ids = []
    if not ids:
        ids = MODELES_PREFERES[:]
    ordre = {m: k for k, m in enumerate(MODELES_PREFERES)}
    return sorted(set(ids), key=lambda m: (ordre.get(m, 99), m))


def appeler_llm(client, modele: str, prompt: str, essais: int = 5) -> dict:
    """Appel chat/completions avec relance automatique (quota, erreurs réseau)."""
    messages = [{"role": "system", "content": PROMPT_SYSTEME},
                {"role": "user", "content": prompt}]
    json_mode = True
    derniere_erreur = None
    for k in range(essais):
        try:
            kwargs = dict(model=modele, messages=messages, temperature=0, max_tokens=16000)
            if json_mode:
                kwargs["response_format"] = {"type": "json_object"}
            r = client.chat.completions.create(**kwargs)
            return parser_json(r.choices[0].message.content or "")
        except Exception as e:  # noqa: BLE001
            derniere_erreur = e
            msg = str(e).lower()
            code = getattr(e, "status_code", None)
            if json_mode and (code in (400, 422) or "response_format" in msg):
                json_mode = False  # le modèle ne supporte pas le mode JSON : on réessaie sans
                continue
            if code in (401, 403):
                raise RuntimeError("Clé API Albert refusée (401/403). Vérifiez la clé.") from e
            if isinstance(e, ValueError) and k < essais - 1:  # JSON illisible -> on relance
                continue
            time.sleep(min(60, 5 * 2 ** k))  # 5, 10, 20, 40 s
    raise RuntimeError(f"Échec de l'appel à Albert après {essais} essais : {derniere_erreur}")


def parser_json(texte: str) -> dict:
    t = texte.strip()
    t = re.sub(r"^```(?:json)?\s*|\s*```$", "", t)
    try:
        return json.loads(t)
    except json.JSONDecodeError:
        m = re.search(r"\{.*\}", t, re.S)
        if m:
            try:
                return json.loads(m.group(0))
            except json.JSONDecodeError:
                pass
    raise ValueError("Réponse non JSON : " + texte[:300])


# ---------------------------------------------------------------------------
# 4. Résultats + contrôles automatiques
# ---------------------------------------------------------------------------
@dataclass
class Resultat:
    ligne: int
    indicateur: str
    valeur: object = None
    confiance: str = "non trouvé"
    page: int | None = None
    extrait: str = ""
    commentaire: str = ""
    controles: list[str] = field(default_factory=list)


def _abaisser(c: str) -> str:
    return {"haute": "moyenne", "moyenne": "faible"}.get(c, c)


def _formes_nombre(x: float) -> list[str]:
    """Écritures possibles d'un nombre dans un PDF français."""
    formes = set()
    for dec in (0, 1, 2, 3):
        s = f"{x:.{dec}f}"
        if dec and float(s) != round(x, dec):
            continue
        ent, _, fr = s.partition(".")
        groupes = f"{int(ent):,}".replace(",", " ") if ent.lstrip("-").isdigit() else ent
        for e in {ent, groupes, groupes.replace(" ", " "), groupes.replace(" ", ".")}:
            formes.add(e + ("," + fr if fr else ""))
            formes.add(e + ("." + fr if fr else ""))
    return [f for f in formes if f]


def valeur_presente(valeur, texte: str) -> bool:
    if valeur is None:
        return True
    t = texte.replace(" ", " ")
    if isinstance(valeur, (int, float)) and not isinstance(valeur, bool):
        return any(re.search(rf"(?<![\d,.]){re.escape(f)}(?![\d])", t)
                   for f in _formes_nombre(float(valeur)))
    v = _norm(valeur)
    if len(v) < 3:
        return True
    tn = _norm(t)
    if v in tn:
        return True
    mots = [m for m in re.findall(r"\w{3,}", v)]
    return bool(mots) and sum(m in tn for m in mots) / len(mots) >= 0.7


def _num(v):
    if isinstance(v, bool) or v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(" ", "").replace(" ", "").replace(",", "."))
    except ValueError:
        return None


def verifier(resultats: dict[int, Resultat], pages: list[str]) -> None:
    """Abaisse la confiance quand une valeur n'est pas retrouvée dans le PDF
    ou quand une règle de cohérence n'est pas respectée."""
    tout = "\n".join(pages)
    for r in resultats.values():
        if r.valeur is None:
            r.confiance = "non trouvé"
            r.extrait, r.page = "", None
            continue
        calcul = bool(re.search(r"calcul|conver|divis|multipli|somm|addition|dedui|estim|ramen", _norm(r.commentaire)))
        page_txt = pages[r.page - 1] if r.page and 1 <= r.page <= len(pages) else ""
        if not valeur_presente(r.valeur, page_txt):
            if valeur_presente(r.valeur, tout):
                r.controles.append("valeur trouvée dans le PDF mais pas à la page indiquée")
                r.confiance = _abaisser(r.confiance) if r.confiance == "haute" else r.confiance
            elif not calcul:
                r.controles.append("valeur introuvable dans le texte du PDF")
                r.confiance = "faible"
            elif r.confiance == "haute":
                r.controles.append("valeur calculée par l'IA (absente telle quelle du PDF)")
                r.confiance = "moyenne"

    par_lib = {_norm(r.indicateur): r for r in resultats.values()}

    def trouve(motif):
        return next((r for k, r in par_lib.items() if re.search(motif, k)), None)

    # SIRET / SIREN / code postal
    for motif, n in ((r"^siret", 14), (r"^siren", 9), (r"^code postal", 5)):
        r = trouve(motif)
        if r and r.valeur is not None:
            chiffres = re.sub(r"\D", "", str(r.valeur))
            if len(chiffres) != n:
                r.controles.append(f"devrait comporter {n} chiffres")
                r.confiance = "faible"
            else:
                r.valeur = chiffres
    # Crelat = 60/50/40 % de Cref (sans modulation)
    cref = trouve(r"^cref")
    vcref = _num(cref.valeur) if cref else None
    if vcref:
        for an, ratio in (("2030", .6), ("2040", .5), ("2050", .4)):
            r = trouve(rf"^crelat {an} sans modulation")
            v = _num(r.valeur) if r else None
            if v and abs(v - ratio * vcref) / (ratio * vcref) > 0.05:
                r.controles.append(f"incohérent avec Cref ({int(ratio*100)} % de {vcref:g} = {ratio*vcref:.1f})")
                r.confiance = _abaisser(r.confiance)
        # Somme des énergies ≈ Cref
        energies = [trouve(m) for m in (r"^electricite - reference", r"^gaz - reference", r"^autres energies - reference")]
        vals = [_num(e.valeur) for e in energies if e and _num(e.valeur) is not None]
        if len(vals) >= 2 and abs(sum(vals) - vcref) / vcref > 0.15:
            for e in energies:
                if e and e.valeur is not None:
                    e.controles.append(f"somme des énergies ({sum(vals):.1f}) ≠ Cref ({vcref:g})")
    # Pourcentages
    for r in resultats.values():
        if "(%)" in r.indicateur:
            v = _num(r.valeur)
            if v is not None and not 0 <= v <= 100:
                r.controles.append("pourcentage hors de 0–100")
                r.confiance = "faible"


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------
def extraire_rapport(client, modele: str, pages: list[str], indicateurs: list[Indicateur],
                     mode: str = "complet", paralleles: int = 2, progression=None
                     ) -> dict[int, Resultat]:
    """Lance un appel par groupe d'indicateurs. mode = 'complet' | 'cible'."""
    groupes: dict[str, list[Indicateur]] = {}
    for i in indicateurs:
        groupes.setdefault(i.groupe, []).append(i)

    resultats = {i.ligne: Resultat(i.ligne, i.libelle) for i in indicateurs}
    erreurs = []

    def tache(g, inds):
        choix = pages_pertinentes(g, inds, pages) if mode == "cible" else None
        return g, appeler_llm(client, modele, construire_prompt(g, inds, pages, choix))

    fait = 0
    with ThreadPoolExecutor(max_workers=max(1, paralleles)) as ex:
        futurs = [ex.submit(tache, g, inds) for g, inds in groupes.items()]
        for f in as_completed(futurs):
            try:
                g, rep = f.result()
                for item in rep.get("resultats", []) if isinstance(rep, dict) else []:
                    try:
                        ligne = int(item.get("id"))
                    except (TypeError, ValueError):
                        continue
                    if ligne not in resultats:
                        continue
                    r = resultats[ligne]
                    val = item.get("valeur")
                    if isinstance(val, str) and _norm(val) in ("", "null", "nc", "n/a", "na", "none", "-"):
                        val = None
                    r.valeur = val
                    conf = _norm(item.get("confiance") or "faible")
                    r.confiance = conf if conf in ("haute", "moyenne", "faible") else "faible"
                    try:
                        r.page = int(item.get("page")) if item.get("page") not in (None, "") else None
                    except (TypeError, ValueError):
                        r.page = None
                    r.extrait = str(item.get("extrait") or "")[:300]
                    r.commentaire = str(item.get("commentaire") or "")[:500]
            except Exception as e:  # noqa: BLE001
                erreurs.append(str(e))
            fait += 1
            if progression:
                progression(fait, len(futurs))
    verifier(resultats, pages)
    if erreurs and all(r.valeur is None for r in resultats.values()):
        raise RuntimeError(" | ".join(dict.fromkeys(erreurs)))
    if erreurs:
        for r in resultats.values():
            if r.valeur is None and not r.commentaire:
                r.commentaire = "Non traité ou erreur API : " + erreurs[0][:200]
    return resultats


# ---------------------------------------------------------------------------
# 5. Écriture Excel
# ---------------------------------------------------------------------------
def _valeur_excel(v):
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return v
    s = str(v).strip()
    n = _num(s)
    if n is not None and re.fullmatch(r"-?[\d\s ]+([.,]\d+)?", s) and not re.fullmatch(r"0\d+", s.replace(" ", "")):
        return int(n) if n.is_integer() and "," not in s and "." not in s else n
    return s


def ecrire_excel(modele_bytes: bytes, feuille: str,
                 rapports: list[tuple[str, dict[int, Resultat]]]) -> bytes:
    """Remplit le modèle : colonne A = indicateurs (inchangée), une colonne par PDF
    à partir de la 1re colonne libre, ligne 1 = en-tête ajoutée."""
    wb = openpyxl.load_workbook(io.BytesIO(modele_bytes))
    ws = wb[feuille]
    ws.insert_rows(1)
    col0 = max(3, ws.max_column + 1)
    gras = Font(bold=True)
    ws.cell(1, 1, "Indicateur").font = gras
    if any(ws.cell(r, 2).value for r in range(2, ws.max_row + 1)):
        ws.cell(1, 2, "Remarque").font = gras
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 60)

    for k, (nom, res) in enumerate(rapports):
        c = col0 + k
        h = ws.cell(1, c, nom)
        h.font = gras
        h.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[get_column_letter(c)].width = 28
        for r in res.values():
            cell = ws.cell(r.ligne + 1, c)  # +1 : ligne d'en-tête insérée
            cell.value = _valeur_excel(r.valeur)
            cell.fill = PatternFill("solid", fgColor=COULEURS.get(r.confiance, "FFFFFF"))
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            note = [f"Confiance : {r.confiance}"]
            if r.page:
                note.append(f"Page {r.page}")
            if r.extrait:
                note.append(f"« {r.extrait} »")
            if r.commentaire:
                note.append(r.commentaire)
            if r.controles:
                note.append("Contrôles : " + " ; ".join(r.controles))
            if r.valeur is not None or r.commentaire:
                cell.comment = Comment("\n".join(note), "Extraction Albert", width=320, height=160)
    ws.freeze_panes = "B2"

    # Feuille détail (format long : une ligne par PDF x indicateur)
    d = wb.create_sheet("Détail extraction")
    entetes = ["Fichier PDF", "Ligne modèle", "Indicateur", "Valeur", "Confiance",
               "Page", "Extrait", "Commentaire", "Contrôles automatiques"]
    d.append(entetes)
    for cell in d[1]:
        cell.font = gras
    for nom, res in rapports:
        for r in res.values():
            d.append([nom, r.ligne, r.indicateur, _valeur_excel(r.valeur), r.confiance,
                      r.page, r.extrait, r.commentaire, " ; ".join(r.controles)])
            d.cell(d.max_row, 5).fill = PatternFill("solid", fgColor=COULEURS.get(r.confiance, "FFFFFF"))
    for col, w in zip("ABCDEFGHI", (30, 8, 50, 22, 12, 6, 50, 40, 40)):
        d.column_dimensions[col].width = w
    d.auto_filter.ref = d.dimensions
    d.freeze_panes = "A2"

    # Légende
    l = wb.create_sheet("Légende")
    l.append(["Couleur", "Signification"])
    for cell in l[1]:
        cell.font = gras
    sens = {
        "haute": "Valeur lue telle quelle dans le rapport et retrouvée dans le texte du PDF",
        "moyenne": "Valeur convertie / calculée, ou correspondance approximative : à vérifier rapidement",
        "faible": "Valeur incertaine ou non retrouvée dans le PDF, ou contrôle de cohérence en échec : à vérifier",
        "non trouvé": "Information absente, anonymisée ou « NC » dans le rapport",
        "validé": "Valeur corrigée ou validée manuellement dans l'application",
    }
    for k, txt in sens.items():
        l.append([k, txt])
        l.cell(l.max_row, 1).fill = PatternFill("solid", fgColor=COULEURS[k])
    l.append([])
    l.append(["", "Survolez une cellule pour voir la page source, l'extrait cité et les contrôles."])
    l.column_dimensions["A"].width = 14
    l.column_dimensions["B"].width = 100

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
